"""Unit tests for document ingestion and classification (core/ingest.py)."""

import io
import os
import tempfile
import pytest
from PIL import Image, ImageDraw
import docx
from reportlab.pdfgen import canvas

from core.ingest import (
    MAX_FILE_SIZE_BYTES,
    classify_document,
    process_documents,
    read_file,
)


@pytest.fixture
def sample_pdf_bytes():
    """Generate in-memory PDF with sample medical lab text."""
    buffer = io.BytesIO()
    c = canvas.Canvas(buffer)
    c.drawString(100, 750, "METROPOLITAN CLINICAL LABORATORY REPORT")
    c.drawString(100, 720, "Patient: John Doe, Age: 34")
    c.drawString(100, 690, "Hemoglobin: 10.2 g/dL (Ref Range: 13.5 - 17.5)")
    c.drawString(100, 660, "ESR: 45 mm/hr (Ref Range: 0 - 15)")
    c.save()
    buffer.seek(0)
    buffer.name = "lab_report.pdf"
    return buffer


@pytest.fixture
def sample_docx_bytes():
    """Generate in-memory DOCX with sample clinical notes."""
    doc = docx.Document()
    doc.add_heading("Hospital Discharge Summary", 0)
    doc.add_paragraph("Admission Date: 2026-09-10. Discharge Date: 2026-09-15.")
    doc.add_paragraph("Chief complaint: Persistent cough and night sweats for 3 weeks.")
    doc.add_paragraph("Assessment and Plan: Advise follow-up with pulmonologist.")
    
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    buffer.name = "discharge_summary.docx"
    return buffer


@pytest.fixture
def sample_image_bytes():
    """Generate in-memory PNG image."""
    img = Image.new("RGB", (300, 100), color=(255, 255, 255))
    draw = ImageDraw.Draw(img)
    draw.text((10, 40), "Rx Amoxicillin 500mg TID 5 days", fill=(0, 0, 0))
    buffer = io.BytesIO()
    img.save(buffer, format="PNG")
    buffer.seek(0)
    buffer.name = "rx_photo.png"
    return buffer


def test_text_pdf_returns_text(sample_pdf_bytes):
    """Acceptance Check 1: A text PDF returns text."""
    res = read_file(sample_pdf_bytes)
    assert res["kind"] == "pdf"
    assert "Hemoglobin" in res["text"]
    assert "ESR" in res["text"]
    assert len(res["warnings"]) == 0


def test_docx_and_txt_return_text(sample_docx_bytes):
    """Acceptance Check 3: A DOCX and a TXT file return text."""
    # Test DOCX
    docx_res = read_file(sample_docx_bytes)
    assert docx_res["kind"] == "docx"
    assert "Discharge Summary" in docx_res["text"]
    assert "Persistent cough" in docx_res["text"]
    assert len(docx_res["warnings"]) == 0

    # Test TXT
    txt_content = b"Patient reports sore throat and mild fever for 2 days. Rx paracetamol 650mg."
    txt_buffer = io.BytesIO(txt_content)
    txt_buffer.name = "notes.txt"
    txt_res = read_file(txt_buffer)
    assert txt_res["kind"] == "txt"
    assert "sore throat" in txt_res["text"]
    assert len(txt_res["warnings"]) == 0


def test_photo_scanned_image_returns_text_or_handled(sample_image_bytes, monkeypatch):
    """Acceptance Check 2: A photo/scanned image returns text."""
    # Mock pytesseract to simulate OCR extraction reliably across platforms
    monkeypatch.setattr(
        "pytesseract.image_to_string",
        lambda img: "Rx: Amoxicillin 500mg capsules. Sig: 1 tab PO TID x 5 days.",
    )
    res = read_file(sample_image_bytes)
    assert res["kind"] == "image"
    assert "Amoxicillin" in res["text"]
    assert len(res["warnings"]) == 0


def test_mixed_5_files_classification():
    """Acceptance Check 4: 5 mixed files are each classified correctly."""
    lab_text = "Metropolitan Labs: Hemoglobin 10.2 g/dL, Ref Range 13.5-17.5. Platelet count 150000. ESR 42 mm/hr."
    assert classify_document(lab_text) == "lab_report"

    rx_text = "Rx: Tab. Amoxicillin 500 mg TID for 5 days. Tab. Paracetamol 650 mg SOS."
    assert classify_document(rx_text) == "prescription"

    discharge_text = "Discharge Summary: Patient admitted on 2026-09-01. Chief complaint: high fever. Assessment and Plan: Recovered."
    assert classify_document(discharge_text) == "doctor_notes_or_discharge_summary"

    imaging_text = "Chest X-Ray PA View: Findings: Apical fibrocavitary infiltration noted. Impression: Consistent with active pulmonary Koch's (TB)."
    assert classify_document(imaging_text) == "imaging_report"

    rash_text = "Symptom Photo Note: Erythematous maculopapular rash observed on both lower extremities with petechiae."
    assert classify_document(rash_text) == "symptom_photo"


def test_empty_or_unreadable_file_returns_warning_not_crash():
    """Acceptance Check 5: An empty or unreadable file returns a warning, not a crash."""
    empty_txt = io.BytesIO(b"")
    empty_txt.name = "empty.txt"
    res = read_file(empty_txt)
    assert res["text"] == ""
    assert any("empty" in w.lower() for w in res["warnings"])

    # Corrupt PDF
    corrupt_pdf = io.BytesIO(b"%PDF-not-a-valid-pdf-header-corrupt")
    corrupt_pdf.name = "corrupt.pdf"
    res_corrupt = read_file(corrupt_pdf)
    assert res_corrupt["text"] == ""
    assert len(res_corrupt["warnings"]) > 0


def test_oversized_file_rejected_with_clear_message():
    """Acceptance Check 6: An oversized file is rejected with a clear message."""
    mock_large_file = io.BytesIO(b"0" * (MAX_FILE_SIZE_BYTES + 1024))
    mock_large_file.name = "huge_scan.pdf"
    res = read_file(mock_large_file)
    assert res["kind"] == "error"
    assert res["text"] == ""
    assert any("exceeds maximum allowed size" in w for w in res["warnings"])


def test_process_documents_batch(sample_pdf_bytes, sample_docx_bytes):
    """Test process_documents with a batch of multiple files."""
    results = process_documents([sample_pdf_bytes, sample_docx_bytes])
    assert len(results) == 2
    assert results[0]["doc_type"] == "lab_report"
    assert results[1]["doc_type"] == "doctor_notes_or_discharge_summary"
